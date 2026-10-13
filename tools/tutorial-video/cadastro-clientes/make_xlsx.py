# Planilha fictícia no formato de um ERP ("Pesquisa de Pessoas") para o tutorial do cadastro de clientes.
import json, openpyxl
from openpyxl.styles import Font
C = json.load(open('clientes/cust.json', encoding='utf-8'))
fmt = lambda d: f'{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}'
disc = lambda v: ';'.join(f'{x:g}%' for x in v)
H = ['Cód.', 'Razão Social', 'Fantasia', 'CPF/CNPJ', 'Cidade', 'UF', 'Vendedor', 'Abat. Financeiro', 'Situação cadastral']
rows = []
for i, c in enumerate(C[:7]):
    r = [c['code'], c['name'], c['trade_name'], fmt(c['cnpj']), c['city'], 'SP', c['seller'], disc(c['disc']), 'Ativo']
    if i == 1: r[7] = '2%;1%'          # desconto mudou
    if i == 3: r[2] = 'Nova Era Itaquá'  # fantasia mudou
    rows.append(r)
rows.append(['00230', 'Mercado Recanto Verde Ltda', 'Recanto Verde', '51.623.847/0001-05', 'Arujá', 'SP', 'Paula Vendas', '3%', 'Ativo'])
rows.append(['00231', 'Recanto Verde Filial Ltda', 'Recanto Verde 2', '51.623.847/0001-05', 'Arujá', 'SP', 'Paula Vendas', '3%', 'Ativo'])
rows.append(['00232', 'Empório Sabor da Serra Ltda', 'Sabor da Serra', '62.734.158/0001-0', 'Mogi das Cruzes', 'SP', 'Roberto Campos', '2%', 'Ativo'])
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Pesquisa de Pessoas'
ws.append(['Pesquisa de Pessoas - Clientes (dados fictícios)']); ws['A1'].font = Font(bold=True, size=13); ws.append([])
ws.append(H)
for c in ws[3]: c.font = Font(bold=True)
for r in rows:
    ws.append(r)
for row in ws.iter_rows(min_row=4):
    row[0].number_format = '@'
for col, w in zip('ABCDEFGHI', [8, 32, 22, 20, 18, 5, 16, 16, 18]): ws.column_dimensions[col].width = w
wb.save('clientes/pesquisa-de-pessoas-ficticio.xlsx'); print(len(rows), 'linhas')
